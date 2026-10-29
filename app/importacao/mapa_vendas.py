# -*- coding: utf-8 -*-
"""Mapa de vendas: a receita bruta faturada por UF do cliente, lida da aba MPP DB do Painel de Resultado.

A aba MPP DB é a base de notas fiscais do Painel (uma linha por produto de cada NF): pedido, cliente, UF, produto, quantidade e
receita bruta (ROB). Ela vem no mesmo arquivo da DRE, então o mapa se alimenta sozinho a cada mês processado, sem planilha nova.

Cada linha do blob: [aaaamm, UF, pedido, cliente, produto, peças, receita bruta]. Linhas iguais são somadas. A UF "EX" é o exterior;
sem UF no cadastro, "ND".
"""
import datetime
import os
import unicodedata

import openpyxl

ABA = 'MPP DB'


def _norm(s):
    return unicodedata.normalize('NFD', str(s or '')).encode('ascii', 'ignore').decode().upper().strip()


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def _ym(dt, ano, mes):
    if isinstance(dt, (datetime.datetime, datetime.date)):
        return dt.year * 100 + dt.month
    if isinstance(ano, (int, float)) and isinstance(mes, (int, float)):
        return int(ano) * 100 + int(mes)
    return None


def ler(painel, ate_ym=None):
    """Blob MAPAVENDAS do Painel, ou None se a aba não existir. `ate_ym` corta o que vem depois do mês processado."""
    wb = openpyxl.load_workbook(painel, read_only=True, data_only=True)
    try:
        if ABA not in wb.sheetnames:
            return None
        linhas = wb[ABA].iter_rows(values_only=True)
        col = None
        for linha in linhas:                       # o cabeçalho é a primeira linha que tem UF e cliente
            nomes = [_norm(c) for c in linha]
            if 'UF' in nomes and 'DESCRICAO CLIENTE' in nomes:
                col = {n: i for i, n in enumerate(nomes) if n}
                break
        if not col:
            return None
        g = lambda linha, nome: linha[col[nome]] if nome in col and col[nome] < len(linha) else None
        soma = {}
        for linha in linhas:
            ym = _ym(g(linha, 'DT EMIS'), g(linha, 'ANO'), g(linha, 'MES'))
            if not ym or (ate_ym and ym > ate_ym):
                continue
            rob = _num(g(linha, 'ROB'))
            if not rob and not _num(g(linha, 'QTDE')):
                continue
            uf = str(g(linha, 'UF') or 'ND').strip().upper() or 'ND'
            ped = g(linha, 'PEDIDO')
            ped = str(int(ped)) if isinstance(ped, float) and ped == int(ped) else str(ped if ped is not None else '')
            chave = (ym, uf, ped, str(g(linha, 'DESCRICAO CLIENTE') or '').strip(), str(g(linha, 'DESC TECNICA') or '').strip())
            o = soma.setdefault(chave, [0.0, 0.0])
            o[0] += _num(g(linha, 'QTDE'))
            o[1] += rob
    finally:
        wb.close()
    if not soma:
        return None
    rows = [[k[0], k[1], k[2], k[3], k[4], round(v[0], 2), round(v[1], 2)] for k, v in sorted(soma.items())]
    return {'fonte': os.path.basename(painel), 'ate': max(r[0] for r in rows), 'rows': rows}
