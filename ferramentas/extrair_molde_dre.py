# -*- coding: utf-8 -*-
"""
Extrai da aba "DRE DINÂMICO" do Painel de Resultado o MOLDE da DRE no padrão da empresa: linhas (rótulo, nível, se
começa recolhida), fórmulas e o estilo de cada célula (fonte, preenchimento, bordas, formato de número, larguras e alturas).
Só o molde — os números vêm das abas REAL/ORÇADO (ver parse_dre_din no Kit).

    python ferramentas/extrair_molde_dre.py "<Painel Resultado.xlsm>" [saida.json] [--cache valores.json]

O molde muda só quando alguém muda o layout da aba; por isso é um arquivo versionado (dre_dinamico_molde.json), não uma
leitura de todo mês: carregar o Painel inteiro com estilos custa memória que o plano da Render não tem.
`--cache` grava também os valores calculados que o Excel deixou na aba (a prova do avaliador de fórmulas).
"""
import colorsys
import json
import os
import re
import sys

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX

COLUNAS = ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N']
PRIMEIRA, ULTIMA = 1, 299          # a aba tem 304 linhas; as últimas estão vazias


def paleta_tema(wb):
    xml = (wb.loaded_theme or b'').decode('utf-8', 'ignore')
    bloco = re.search(r'<a:clrScheme[^>]*>(.*?)</a:clrScheme>', xml, re.S)
    ordem = ['dk1', 'lt1', 'dk2', 'lt2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6']
    cores = {}
    for nome in ordem:
        m = re.search(r'<a:%s>\s*<a:(?:srgbClr val|sysClr val="[^"]*" lastClr)="([0-9A-Fa-f]{6})"' % nome, bloco.group(1) if bloco else '')
        cores[nome] = m.group(1).upper() if m else None
    # índice de tema do Excel: 0=lt1 1=dk1 2=lt2 3=dk2 4..9=accent1..6
    return [cores['lt1'], cores['dk1'], cores['lt2'], cores['dk2']] + [cores['accent%d' % i] for i in range(1, 7)]


def com_tint(hexa, tint):
    if not tint:
        return hexa
    r, g, b = (int(hexa[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return '%02X%02X%02X' % tuple(round(x * 255) for x in (r, g, b))


def resolve(cor, tema, padrao=None):
    if cor is None:
        return padrao
    if cor.type == 'rgb' and isinstance(cor.rgb, str):
        return cor.rgb[-6:].upper()
    if cor.type == 'theme' and cor.theme is not None and cor.theme < len(tema) and tema[cor.theme]:
        return com_tint(tema[cor.theme], cor.tint or 0)
    if cor.type == 'indexed' and cor.indexed is not None and cor.indexed < len(COLOR_INDEX):
        return str(COLOR_INDEX[cor.indexed])[-6:].upper()
    return padrao


def estilo_da(c, tema):
    f, fl, b, al = c.font, c.fill, c.border, c.alignment
    lado = lambda s: ([s.style, resolve(s.color, tema, '000000')] if s is not None and s.style else None)
    preench = None
    if fl is not None and fl.fill_type:
        preench = {'t': fl.fill_type, 'fg': resolve(fl.fgColor, tema, '000000'), 'bg': resolve(fl.bgColor, tema, 'FFFFFF')}
    est = {'fn': f.name, 'sz': f.sz, 'b': bool(f.b), 'i': bool(f.i), 'u': bool(f.u),
           'co': resolve(f.color, tema, '000000'), 'fill': preench, 'nf': c.number_format, 'h': al.horizontal,
           'v': al.vertical, 'ind': al.indent or 0, 'wrap': bool(al.wrap_text),
           'bd': {k: lado(getattr(b, n)) for k, n in (('t', 'top'), ('b', 'bottom'), ('l', 'left'), ('r', 'right'))}}
    est['bd'] = {k: v for k, v in est['bd'].items() if v}
    return {k: v for k, v in est.items() if v not in (None, False, 0, '', {})}


def normaliza(formula):
    f = formula.lstrip('=').replace('$', '')
    f = re.sub(r"'DE\|PARA'!I3", 'MES', f)
    f = re.sub(r"'DE\|PARA'!W3", 'FLAG', f)
    return f


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    cache = sys.argv[sys.argv.index('--cache') + 1] if '--cache' in sys.argv else None
    if cache and cache in args:
        args.remove(cache)
    origem = args[0]
    saida = args[1] if len(args) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                                      'app', 'importacao', 'dre_dinamico_molde.json')
    wb = openpyxl.load_workbook(origem, read_only=False, keep_vba=False)
    wv = openpyxl.load_workbook(origem, read_only=False, data_only=True, keep_vba=False) if cache else None
    ws = wb['DRE DINÂMICO']
    tema = paleta_tema(wb)
    estilos, idx = [], {}

    def estilo_id(c):
        e = estilo_da(c, tema)
        chave = json.dumps(e, sort_keys=True, ensure_ascii=False)
        if chave not in idx:
            idx[chave] = len(estilos)
            estilos.append(e)
        return idx[chave]

    cols = {}
    for letra, d in ws.column_dimensions.items():          # larguras (cada entrada cobre de min a max)
        for n in range(d.min, min(d.max, 30) + 1):
            cols[openpyxl.utils.get_column_letter(n)] = round(d.width or 8.43, 2)
    linhas = []
    for r in range(PRIMEIRA, ULTIMA + 1):
        d = ws.row_dimensions[r]
        cel = {}
        for col in COLUNAS:
            c = ws['%s%d' % (col, r)]
            item = {'s': estilo_id(c)}
            v = c.value
            if isinstance(v, str) and v.startswith('='):
                item['f'] = 'L' if 'INDIRECT' in v else normaliza(v)
            elif v is not None:
                item['v'] = v
            cel[col] = item
        a = ws['A%d' % r].value
        if isinstance(a, str) and a.startswith('='):
            a = normaliza(a)
            a = {'f': a}
        linha = {'r': r, 'h': d.height, 'n': d.outlineLevel or 0, 'o': bool(d.hidden), 'cel': cel}
        if a not in (None, ''):
            linha['a'] = a
        linhas.append(linha)
    flag = None
    if wv is not None:
        flag = wv['DE|PARA']['W3'].value
    molde = {'versao': 1, 'origem': os.path.basename(origem), 'zoom': ws.sheet_view.zoomScale, 'flag': flag,
             'colunas': {k: v for k, v in cols.items() if k in COLUNAS}, 'estilos': estilos, 'linhas': linhas}
    os.makedirs(os.path.dirname(saida), exist_ok=True)
    json.dump(molde, open(saida, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print('molde: %d linhas, %d estilos, flag=%r -> %s (%d KB)' % (len(linhas), len(estilos), flag, saida, os.path.getsize(saida) // 1024))
    if cache:
        wv_ws = wv['DRE DINÂMICO']
        valores = {}
        for r in range(PRIMEIRA, ULTIMA + 1):
            valores[r] = {col: wv_ws['%s%d' % (col, r)].value for col in ('A', 'C', 'E', 'F', 'G', 'I', 'K', 'L', 'M')
                          if wv_ws['%s%d' % (col, r)].value is not None}
        json.dump({'mes': wv['DE|PARA']['I3'].value, 'unidade': wv_ws['A2'].value, 'ano': wv_ws['E9'].value, 'valores': valores},
                  open(cache, 'w', encoding='utf-8'), ensure_ascii=False, default=str)
        print('valores em cache gravados em', cache)


if __name__ == '__main__':
    main()
